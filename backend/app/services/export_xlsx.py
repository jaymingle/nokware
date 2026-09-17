"""An Ask answer as an Excel workbook: the live report figures as a sheet you can work with, and its chart.

The chart is a real Excel chart over the Figures cells, not a picture, so changing a cell moves the bar.

- "Fewer than 5" is never a number. Its count cell stays empty and the text goes in "Shown as", so nobody can total
  a column and reveal a suppressed count. The same for a figure that is a range.
- Only live report figures and budget figures (each block proved against its stated total) become cells. Numbers
  quoted from other documents stay in the answer's text: a sheet invites sums across rows the Ledger can't support.

A cell that would start a formula (=, +, -, @) is prefixed with an apostrophe:
the question is the resident's own words.
"""

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.schemas.ask import AskChart, AskFigure
from app.services.ask_export import HEADER_NOTICE, Content, chart_footnote, figure_footnote, figures_heading, figures_notes
from app.services.export_csv import answer_text, figure_value

INK = "FF17242B"
TEAL = "FF1F6F5C"
HEADING = Font(bold=True, color=INK)
WHITE_HEADING = Font(bold=True, color="FFFFFFFF")
SMALL = Font(size=9, color="FF4A5A5F")
FILL = PatternFill("solid", fgColor=TEAL)
WRAP = Alignment(wrap_text=True, vertical="top")
_FORMULA = ("=", "+", "-", "@", "\t", "\r")
MAX_ROWS_CHARTED = 24  # a chart of more categories than this is unreadable; the cells still hold every row


def _safe(value: object) -> object:
    return f"'{value}" if isinstance(value, str) and str(value).startswith(_FORMULA) else value


def _write(sheet: Worksheet, row: int, values: list[object], font: Font | None = None) -> int:
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=column, value=_safe(value))
        if font:
            cell.font = font
    return row + 1


def _widths(sheet: Worksheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _answer_sheet(sheet: Worksheet, content: Content) -> None:
    sheet.sheet_properties.tabColor = TEAL
    _widths(sheet, [22, 70, 28, 30])
    row = _write(sheet, 1, [HEADER_NOTICE], SMALL)
    row = _write(sheet, row + 1, ["Question", content.question], HEADING)
    row = _write(sheet, row, ["Answered", content.answered])
    row += 1
    sheet.cell(row=row, column=1, value="Answer").font = HEADING
    answer = sheet.cell(row=row, column=2, value=_safe(answer_text(content)))
    answer.alignment = WRAP
    sheet.row_dimensions[row].height = 220
    row += 2
    for note in (content.no_information, content.chart_note, content.chart.note if content.chart else None):
        if note:
            row = _write(sheet, row, ["Note", note], SMALL)
    if content.sources:
        row = _write(sheet, row + 1, ["Sources"], HEADING)
        row = _write(sheet, row, ["Cited", "Document", "Department", "Year", "Published at", "Nokware's copy"], HEADING)
        for number, source, provenance in content.sources:
            row = _write(sheet, row, [f"[S{number}]", source.title, source.department_name, source.document_year,
                                      source.source_url, source.ledger_url])
            if provenance:
                row = _write(sheet, row, ["", provenance], SMALL)


def _figure_block(sheet: Worksheet, row: int, number: int, figure: AskFigure) -> tuple[int, tuple[int, int] | None]:
    """Returns the next row and the rows charted."""
    row = _write(sheet, row, [f"F{number}", figure.description], HEADING)
    row = _write(sheet, row, ["Source", figure_footnote(figure)], SMALL)
    header = row
    measure = "Amount (GH¢)" if figure.source == "documents" else "Count"
    for column, title in enumerate(["Category", measure, "Shown as"], start=1):
        cell = sheet.cell(row=header, column=column, value=title)
        cell.font, cell.fill = WHITE_HEADING, FILL
    row = _write(sheet, header + 1, ["Total", figure_value(figure.value), figure.value])
    first = row
    for line in figure.rows:
        row = _write(sheet, row, [line.name, figure_value(line.value), line.value])
    return row + 1, (first, row - 1) if figure.rows else None


def _figures_sheet(sheet: Worksheet, content: Content) -> dict[str, tuple[int, int]]:
    """Returns each figure's breakdown rows, by label, for the chart to point at."""
    sheet.freeze_panes = "A2"
    _widths(sheet, [44, 12, 18])
    row = _write(sheet, 1, [figures_heading(content.figures)], HEADING)
    for note in figures_notes(content.figures):
        row = _write(sheet, row, [note], SMALL)
    row += 1
    ranges: dict[str, tuple[int, int]] = {}
    for number, figure in content.figures:
        row, rows = _figure_block(sheet, row, number, figure)
        if rows:
            ranges[figure.label] = rows
    return ranges


def _chart(chart: AskChart) -> BarChart | LineChart | PieChart:
    """The kind Nokware drew, so the workbook says what the page said."""
    if chart.kind == "line":
        return LineChart()
    if chart.kind in ("pie", "donut"):
        return PieChart()
    drawn = BarChart()
    drawn.type = "bar" if chart.horizontal else "col"
    if chart.kind == "stacked_bar":
        drawn.grouping, drawn.overlap = "stacked", 100
    return drawn


def _chart_sheet(sheet: Worksheet, figures: Worksheet, chart: AskChart, charted: list[tuple[int, int]]) -> None:
    # The bars the page drew and no more, so the workbook's chart matches the note beside it ("the 20 largest of 42").
    drawn_rows = min(len(chart.categories), MAX_ROWS_CHARTED)
    first, last = charted[0][0], min(charted[0][1], charted[0][0] + drawn_rows - 1)
    drawn = _chart(chart)
    drawn.title = chart.title
    drawn.legend = None
    drawn.height, drawn.width = 10, 20
    drawn.add_data(Reference(figures, min_col=2, min_row=first, max_row=last))
    drawn.set_categories(Reference(figures, min_col=1, min_row=first, max_row=last))
    sheet.add_chart(drawn, "A4")
    _widths(sheet, [110])
    row = _write(sheet, 1, [chart.title], HEADING)
    _write(sheet, row, [chart_footnote(chart)], SMALL)


def xlsx(content: Content) -> bytes:
    """Only an answer with live figures is worth one."""
    book = Workbook()
    book.properties.creator = "Nokware (not an official AMA document)"
    book.properties.title = content.question[:255]
    answer: Any = book.active
    answer.title = "Answer"
    _answer_sheet(answer, content)
    figures = book.create_sheet("Figures")
    ranges = _figures_sheet(figures, content)
    charted = [ranges[label] for label in content.chart.figures if label in ranges] if content.chart else []
    if content.chart and charted:
        _chart_sheet(book.create_sheet("Chart"), figures, content.chart, charted)
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()
