"""Ask's charts (only live figures, never a false value) and exports (PDF, Word, CSV: signed, and never mistaken for AMA's)."""

import csv
import io
import json
from collections import defaultdict, deque
from datetime import datetime, timezone
from typing import Any

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader

from app.main import app
from app.routes import ask as ask_route
from app.services import ask_charts, rag, rate_limit, stats
from app.services.ask_charts import DOCUMENT_CHART_REFUSAL, chart_for
from app.services.ask_document_charts import Plotted
from app.services.ask_export import FOOTER_NOTICE, HEADER_NOTICE, Answered, content, export_view, filename, provenance_line, verified
from app.services.ask_figures import CountReports, FigurePlan, count_figure
from app.services.export_csv import csv_bytes
from app.services.export_docx import docx
from app.services.export_pdf import pdf
from app.services.retrieval import Chunk, Retrieval, RetrievedChunk

NOW = datetime(2026, 9, 14, 22, 46, tzinfo=timezone.utc)
AT = "2026-09-14T22:40:00+00:00"


def figure(label: str, description: str, value: str, rows: list[tuple[str, str]] = (), by: str = "none", cited: bool = True) -> dict[str, Any]:
    return {"label": label, "cited": cited, "description": description, "value": value,
            "rows": [{"name": n, "value": v} for n, v in rows], "counted_at": AT, "grouped_by": by if rows else "none"}


TOPICS = figure("R1", "Reports · since Nokware began", "47", [("Solid waste and dumping", "23"), ("Street lighting", "12"),
                                                             ("Drainage and flooding", "fewer than 5")], "topic")
OPEN = figure("R1", "Open reports · since Nokware began", "30", [("Okaikoi South", "14"), ("Ablekuma South", "9"), ("Ashiedu Keteke", "7")], "sub_metro")
RESOLVED = figure("R2", "Resolved reports · since Nokware began", "18", [("Okaikoi South", "6"), ("Ablekuma South", "7"), ("Ashiedu Keteke", "5")], "sub_metro")
RESOLVED_SMALL = figure("R2", "Resolved reports · since Nokware began", "12", [("Okaikoi South", "6"), ("Ablekuma South", "fewer than 5"),
                                                                              ("Ashiedu Keteke", "5")], "sub_metro")
MONTHS = figure("R1", "Reports · this year", "40", [("Jul 2026", "fewer than 5"), ("Aug 2026", "12"), ("Sep 2026", "25")], "month")


# Charts: the kind asked for, unless it can't show the data honestly.

@pytest.mark.parametrize(("question", "figures", "kind", "horizontal", "why"), [
    ("Show open reports by sub-metro as a chart", [OPEN], "bar", False, None),  # unnamed: bars to compare
    ("Graph reports by month this year", [MONTHS], "line", False, None),  # unnamed: a line over time
    ("Chart reports by topic", [TOPICS], "bar", True, None),  # long names lie flat
    ("A pie chart of open reports by sub-metro", [OPEN], "pie", False, None),
    ("A donut chart of open reports by sub-metro", [OPEN], "donut", False, None),
    ("A horizontal bar chart of open reports by sub-metro", [OPEN], "bar", True, None),
    ("A stacked bar of open and resolved by sub-metro", [OPEN, RESOLVED], "stacked_bar", False, None),
    ("A pie chart of reports by topic", [TOPICS], "bar", True, "without inventing a slice size"),  # a "fewer than 5" slice
    ("A pie of open and resolved by sub-metro", [OPEN, RESOLVED], "bar", False, "holds one set of numbers"),
    ("A pie of open versus resolved", [figure("R1", "Open reports", "30"), figure("R2", "Resolved reports", "18")], "bar", False,
     "parts of one whole"),
    ("A stacked bar of open and resolved by sub-metro", [OPEN, RESOLVED_SMALL], "bar", False, "shift every segment above it"),
    ("A stacked bar chart of open reports by sub-metro", [OPEN], "bar", False, "nothing to stack"),
    ("A line graph of open reports by sub-metro", [OPEN], "bar", False, "suggest a trend"),
    ("A scatter plot of reports this year", [MONTHS], "line", False, "doesn't fit counts like these"),
    ("A line chart of reports by month", [MONTHS], "line", False, None),  # a line may carry a range: drawn dashed
])
def test_the_chart_is_the_kind_asked_for_unless_it_would_be_dishonest(
    question: str, figures: list[dict[str, Any]], kind: str, horizontal: bool, why: str | None
) -> None:
    chart, note = chart_for(question, figures)
    assert chart is not None and note is None
    assert (chart["kind"], chart["horizontal"]) == (kind, horizontal)
    assert (why is None and chart["note"] is None) or (why is not None and why in chart["note"])


def test_suppression_is_a_range_never_a_value_and_the_axis_is_round() -> None:
    chart, _ = chart_for("chart open and resolved by sub-metro", [OPEN, RESOLVED_SMALL])
    small = chart["series"][1]["values"][1]
    assert small == {"shown": "fewer than 5", "low": 1, "high": 4}
    assert (chart["axis_max"], chart["ticks"]) == (15, [0, 5, 10, 15]) and chart["figures"] == ["R1", "R2"]
    assert [s["name"] for s in chart["series"]] == ["Open reports", "Resolved reports"] and chart["title"] == "Reports · since Nokware began"
    stacked, _ = chart_for("stacked bar of open and resolved by sub-metro", [OPEN, RESOLVED])
    assert stacked["axis_max"] == 20  # above the tallest stack (14 + 6)


@pytest.mark.parametrize(("question", "figures", "note"), [
    ("What does AMA charge for a stall?", [OPEN], None),  # no chart asked for
    ("Chart the open reports", [figure("R1", "Open reports", "30")], ask_charts.ONE_COUNT),
    ("Graph reports each month", [figure("R1", "Reports", "12", [("Sep 2026", "12")], "month")], ask_charts.ONE_MONTH),
    ("Chart the open reports by sub-metro", [figure("R1", "Open reports", "none", [("Okaikoi South", "none"), ("Kinka", "none")], "sub_metro")], ask_charts.ALL_ZERO),
])
def test_no_chart_when_none_is_asked_for_or_there_is_nothing_to_chart(question: str, figures: list[dict[str, Any]], note: str | None) -> None:
    assert chart_for(question, figures) == (None, note)


def _answer(monkeypatch: pytest.MonkeyPatch, question: str, figures: FigurePlan, text: str,
            plotted: Plotted | None = None, passage: str = "Budget text") -> dict[str, Any]:
    class Model:
        def invoke(self, prompt_input: dict[str, str]) -> str:
            return text

    chunk = RetrievedChunk(chunk=Chunk(1, "d1", 0, passage), score=1.0, document={"title": "2026 Budget", "department": "dept-finance"})
    monkeypatch.setattr(rag, "figures_to_chart", lambda q, a, passages: plotted)  # the reading is tested on its own
    monkeypatch.setattr(rag, "retrieve", lambda q: Retrieval(queries=[q], chunks=[chunk]))
    monkeypatch.setattr(rag, "plan_figures", lambda q, now: figures)
    monkeypatch.setattr(rag, "_answer_chain", lambda: Model())
    return dict(rag.answer_question(question))


def test_a_chart_of_document_data_is_refused_in_plain_words_and_the_figures_stay_in_text(monkeypatch: pytest.MonkeyPatch) -> None:
    result = _answer(monkeypatch, "Give me a bar graph of the 2026 budget by department", rag.NO_FIGURES, "Works gets GH¢2m [S1].")
    assert result["answer"] == f"{DOCUMENT_CHART_REFUSAL}\n\nWorks gets GH¢2m [S1]." and result["chart"] is None
    result = _answer(monkeypatch, "Give me a bar graph of the 2026 budget", rag.NO_FIGURES, "I don't have information on that in the Ledger.")
    assert result["status"] == "no_information" and DOCUMENT_CHART_REFUSAL not in result["answer"]


def test_document_figures_are_charted_when_each_one_is_proved_against_a_cited_passage(monkeypatch: pytest.MonkeyPatch) -> None:
    plotted = Plotted("Market stall fees, 2026 (GHS)", [("Stores - A", "800.00", 800.0), ("Stores - B", "200.00", 200.0)])
    result = _answer(monkeypatch, "Show the stall fees as a bar chart", rag.NO_FIGURES, "Stores A cost 800.00 [S1].", plotted)
    chart = result["chart"]
    assert chart["source"] == "documents" and chart["kind"] == "bar" and chart["counted_at"] is None and chart["figures"] == []
    assert chart["categories"] == ["Stores - A", "Stores - B"] and [v["shown"] for v in chart["series"][0]["values"]] == ["800.00", "200.00"]
    assert chart["axis_max"] == 800 and chart["ticks"] == [0, 200, 400, 600, 800]
    assert DOCUMENT_CHART_REFUSAL not in result["answer"]


def test_a_chart_of_live_figures_comes_with_the_answer(monkeypatch: pytest.MonkeyPatch) -> None:
    cases = [{"category": "civic_service", "isSensitive": False, "topic": "solid_waste", "subMetro": sub, "status": "assigned",
              "recipients": [], "wardLocation": None, "createdAt": "2026-09-10T09:00:00+00:00"} for sub in ["okaikoi-south"] * 6 + ["ablekuma-south"] * 2]
    counted = count_figure(CountReports(group_by="sub_metro"), "R1", cases, NOW, AT)
    result = _answer(monkeypatch, "Show reports by sub-metro as a pie chart", FigurePlan([counted], False), "Okaikoi South leads [R1].")
    assert result["chart"]["kind"] == "bar" and "fewer than 5" in result["chart"]["note"]  # Ablekuma South's 2 is suppressed
    assert DOCUMENT_CHART_REFUSAL not in result["answer"]


def test_a_breakdown_by_month_runs_oldest_first_with_the_empty_months() -> None:
    cases = [{"category": "civic_service", "isSensitive": False, "topic": "solid_waste", "status": "assigned", "recipients": [],
              "createdAt": created} for created in ("2026-07-02T09:00:00+00:00", *["2026-09-03T09:00:00+00:00"] * 6)]
    months = stats.by_month(cases, stats.ReportFilter(), NOW)
    assert months == [("2026-07", 1), ("2026-08", 0), ("2026-09", 6)]
    counted = count_figure(CountReports(group_by="month"), "R1", cases, NOW, AT)
    assert counted.rows == [("Jul 2026", "fewer than 5"), ("Aug 2026", "none"), ("Sep 2026", "6")] and counted.grouped_by == "month"
    this_year = stats.by_month(cases, stats.ReportFilter(period=stats.Period.THIS_YEAR), NOW)
    assert [m for m, _ in this_year] == ["2026-07", "2026-08", "2026-09"]  # nothing before Nokware's first report
    waste_only = stats.by_month([*cases, {**cases[0], "topic": "roads", "createdAt": "2026-05-01T09:00:00+00:00"}],
                                stats.ReportFilter(topic="solid_waste"), NOW)
    assert waste_only[0] == ("2026-05", 0)  # from Nokware's first report of any kind: May really had none


# Exports: signed, laid out once, and every page saying what it is.

def _source(label: str, document_id: str, title: str, cited: bool = True, **extra: Any) -> dict[str, Any]:
    return {"label": label, "cited": cited, "document_id": document_id, "title": title, "chunk_text": "…", "department": None,
            "department_name": "Budget & Rating", "source_type": None, "provenance": "ama_website",
            "source_url": f"https://ama.gov.gh/{document_id}.pdf", "published_at": None, "document_year": 2026, **extra}


QUESTION = "=HYPERLINK(\"x\") What does AMA charge for a stall, as a pie chart by sub-metro?"
ANSWER = ("Stalls cost **GH₵30.00** a month [S2], set by resolution [S1] [S9].\n\n* Lock-up: GH¢50.00 [S2]\n* Open: GH₵30.00 [S2]\n\n"
          "Residents filed 30 open reports [R1]; traders say \"ɛyɛ\".")


@pytest.fixture
def view(monkeypatch: pytest.MonkeyPatch) -> Any:
    settings = ask_route.get_settings().model_copy(update={"public_api_url": "https://api.nokware.example"})
    monkeypatch.setattr("app.services.ask_export.get_settings", lambda: settings)
    chart, note = chart_for(QUESTION, [OPEN])
    sources = [_source("S1", "d1", "AMA Bye-Laws"), _source("S2", "d2", "2026 Fee-Fixing Resolution"), _source("S2", "d2", "2026 Fee-Fixing Resolution"),
               _source("S3", "d3", "Uncited plan", cited=False)]
    return export_view(Answered(QUESTION, ANSWER, "answered", sources, [OPEN, {**RESOLVED, "cited": False}], chart, note), NOW)


def test_the_view_carries_only_what_was_cited_and_any_change_breaks_its_signature(view: Any) -> None:
    assert [s.label for s in view.sources] == ["S1", "S2"] and [f.label for f in view.figures] == ["R1"]
    assert view.sources[1].ledger_url == "https://api.nokware.example/api/ledger/d2/file"
    assert verified(view)
    for change in ({"answer": view.answer + " And more."}, {"question": "Something else?"}, {"answered_at": "2027-01-01T00:00:00+00:00"}):
        assert not verified(view.model_copy(update=change))
    tampered = view.model_copy(update={"sources": [view.sources[0].model_copy(update={"title": "Official AMA statement"})]})
    assert not verified(tampered)


def test_citations_are_renumbered_in_order_and_provenance_reads_as_on_the_web(view: Any) -> None:
    laid_out = content(view, "https://nokware.example")
    paragraph = "".join(text for text, _ in laid_out.blocks[0].runs)
    assert paragraph == "Stalls cost GH₵30.00 a month [1], set by resolution [2] ."  # [S9] was never a source: dropped
    assert [(n, s.title) for n, s, _ in laid_out.sources] == [(1, "2026 Fee-Fixing Resolution"), (2, "AMA Bye-Laws")]
    assert [b.kind for b in laid_out.blocks] == ["paragraph", "bullet", "bullet", "paragraph"] and laid_out.figures[0][0] == 1
    src = view.sources[0]
    assert provenance_line(src) == "Published by Budget & Rating on ama.gov.gh"
    assert provenance_line(src.model_copy(update={"provenance": "department_portal"})) == "Submitted by Budget & Rating · Source: ama.gov.gh"
    assert provenance_line(src.model_copy(update={"provenance": "contributor", "source_url": None})) == "Verified contributor"
    assert filename(view, "pdf") == "nokware-answer-2026-09-14-hyperlink-x-what-does-ama-charge.pdf"


def test_the_pdf_says_on_every_page_it_is_nokwares_not_amas_and_sets_the_cedi(view: Any) -> None:
    reader = PdfReader(io.BytesIO(pdf(content(view, "https://nokware.example"))))
    pages = [page.extract_text() for page in reader.pages]
    for number, text in enumerate(pages, 1):
        assert "not an official AMA document" in text and "Machine-written by Nokware" in text and f"Page {number} of {len(pages)}" in text
    whole = " ".join(pages)
    assert "GH₵30.00" in whole and "ɛyɛ" in whole and "Published by Budget & Rating on ama.gov.gh" in whole
    assert "Open reports · since Nokware began" in whole  # the chart's title, over its image
    links = [a.get_object()["/A"]["/URI"] for p in reader.pages for a in (p.get("/Annots") or [])]
    assert "https://ama.gov.gh/d2.pdf" in links and "https://api.nokware.example/api/ledger/d1/file" in links
    assert reader.metadata.author == "Nokware (not an official AMA document)"


def test_the_word_file_has_the_notices_on_every_page_links_the_chart_and_figures(view: Any) -> None:
    document = Document(io.BytesIO(docx(content(view, "https://nokware.example"))))
    header = " ".join(p.text for p in document.sections[0].header.paragraphs)
    assert "Nokware" in header and HEADER_NOTICE in header and FOOTER_NOTICE in document.sections[0].footer.paragraphs[0].text
    styles = [p.style.name for p in document.paragraphs]
    assert styles.count("List Bullet") == 2 and "Title" in styles
    assert len(document.inline_shapes) == 1 and len(document.tables) == 1
    targets = [rel.target_ref for rel in document.part.rels.values() if rel.reltype.endswith("/hyperlink")]
    assert "https://ama.gov.gh/d1.pdf" in targets and "https://api.nokware.example/api/ledger/d2/file" in targets
    assert document.core_properties.author == "Nokware (not an official AMA document)"


def test_the_csv_holds_figures_as_rows_keeps_suppression_out_of_sums_and_defuses_formulas(view: Any) -> None:
    small = view.model_copy(update={"figures": [view.figures[0].model_copy(update={"rows": [
        view.figures[0].rows[0], view.figures[0].rows[1].model_copy(update={"value": "fewer than 5"})]})]})
    raw = csv_bytes(content(small, "https://nokware.example"))
    assert raw.startswith("﻿".encode())
    rows = list(csv.DictReader(io.StringIO(raw.decode("utf-8-sig"))))
    assert rows[0]["section"] == "notice" and "not an official AMA document" in rows[0]["item"]
    assert rows[1]["item"].startswith("'=HYPERLINK")  # the resident's words can't run as a formula
    figures = [r for r in rows if r["section"] == "figure"]
    assert [(r["category"], r["value"], r["shown_as"]) for r in figures] == [("Total", "30", "30"), ("Okaikoi South", "14", "14"),
                                                                             ("Ablekuma South", "", "fewer than 5")]
    assert [r["item"] for r in rows if r["section"] == "source"] == ["2026 Fee-Fixing Resolution", "AMA Bye-Laws"]
    assert "GH₵30.00" in next(r["item"] for r in rows if r["section"] == "answer")


def test_the_route_renders_only_a_signed_view_as_an_attachment(view: Any, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(rate_limit.EXPORTS, "_hits", defaultdict(deque))
    client = TestClient(app)
    body = json.loads(view.model_dump_json())
    for form, media in (("pdf", "application/pdf"), ("docx", "application/vnd.openxmlformats"), ("csv", "text/csv")):
        response = client.post("/api/ask/export", json={"view": body, "format": form})
        assert response.status_code == 200 and response.headers["content-type"].startswith(media)
        assert response.headers["content-disposition"] == f'attachment; filename="{filename(view, form)}"'
    forged = client.post("/api/ask/export", json={"view": {**body, "answer": "The MCE spent GH₵1bn on a palace."}, "format": "pdf"})
    assert forged.status_code == 403


def test_exports_are_rate_limited(view: Any, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(rate_limit.EXPORTS, "limit", 1)
    monkeypatch.setattr(rate_limit.EXPORTS, "_hits", defaultdict(deque))
    client = TestClient(app)
    body = {"view": json.loads(view.model_dump_json()), "format": "csv"}
    assert [client.post("/api/ask/export", json=body).status_code for _ in range(2)] == [200, 429]


def test_both_ask_routes_send_the_signed_view(monkeypatch: pytest.MonkeyPatch) -> None:
    chunk = RetrievedChunk(chunk=Chunk(1, "d1", 0, "Fee text"), score=1.0, document={"title": "Fees", "department": "dept-finance"})

    class Model:
        def invoke(self, prompt_input: dict[str, str]) -> str:
            return "Stalls cost GH¢30 [S1]."

        def stream(self, prompt_input: dict[str, str]) -> Any:
            yield "Stalls cost GH¢30 [S1]."

    monkeypatch.setattr(rag, "retrieve", lambda q: Retrieval(queries=[q], chunks=[chunk]))
    monkeypatch.setattr(rag, "plan_figures", lambda q, now: rag.NO_FIGURES)
    monkeypatch.setattr(rag, "_answer_chain", lambda: Model())
    client = TestClient(app)
    whole = client.post("/api/ask", json={"question": "What do stalls cost?"}).json()
    done = json.loads(client.post("/api/ask/stream", json={"question": "What do stalls cost?"}).text.splitlines()[-1])
    for export in (whole["export"], done["export"]):
        assert export["question"] == "What do stalls cost?" and [s["title"] for s in export["sources"]] == ["Fees"]
        assert verified(ask_route.ExportView.model_validate(export))
