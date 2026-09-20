"""A budget figure is cited, stripped, exported and described correctly on every path an answer takes.

B was added to the page and to sanitize_citations when budget figures arrived, and missed in five other places.
So a budget answer showed "[B1]" raw in WhatsApp, SMS and every export, read it aloud as "B one", and — worst —
every export and every chat answer called its amounts "counts of the reports residents filed with Nokware, not
figures from a published document", which is exactly backwards. None of those paths had a budget test.
"""

import csv
import io
import re
from typing import Any

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from app.services import ask_export, ask_language, channel_answers, voice_speech
from app.services.ask_export import BUDGET_DATA_NOTE, LIVE_DATA_NOTE, Answered, content, export_view
from app.services.citations import FIGURE_KINDS, KINDS
from app.services.export_csv import csv_bytes, figure_value
from app.services.export_docx import docx
from app.services.export_pdf import pdf
from app.services.export_xlsx import xlsx
from tests.test_ask_export import NOW

SITE = "https://nokware.example"
RAW = re.compile(r"\[[SRB]\d+\]")
ROWS = [("Head Office", "GH¢ 20,270,110"), ("Public Works", "GH¢ 20,232,848"), ("Education", "GH¢ 7,786,630")]
BUDGET = {"label": "B1", "cited": True, "description": "Approved budget · 2026 · by department", "value": "GH¢ 48,289,588",
          "rows": [{"name": name, "value": value} for name, value in ROWS], "grouped_by": "department", "counted_at": None,
          "source": "documents", "document_id": "ama-2026", "document_title": "2026 AMA Budget", "year": 2026, "coverage": "96%"}
CHART = {"kind": "bar", "horizontal": True, "title": "Approved budget · 2026 · by department", "categories": [n for n, _ in ROWS],
         "series": [{"name": "Figures", "values": [{"shown": v, "low": float(v[4:].replace(",", "")), "high": float(v[4:].replace(",", ""))} for _, v in ROWS]}],
         "over_time": False, "figures": ["B1"], "counted_at": None, "source": "documents", "axis_max": 25_000_000,
         "ticks": [0, 5_000_000, 10_000_000, 15_000_000, 20_000_000, 25_000_000], "note": None}
ANSWER = "AMA approved **GH¢ 20,270,110** for Head Office and **GH¢ 20,232,848** for Public Works in 2026 [B1]."


def _content(chart: dict[str, Any] | None = CHART) -> Any:
    return content(export_view(Answered("How much did each department get in 2026? As a chart.", ANSWER, "answered", [], [BUDGET], chart, None), NOW), SITE)


def test_every_citation_kind_is_defined_once_and_used_everywhere() -> None:
    assert KINDS == "SRB" and FIGURE_KINDS == "RB"
    for pattern in (ask_export._CITATION, ask_language._LABEL, voice_speech._TAG, channel_answers._ANY_TAG):
        assert pattern.search("[B1]"), pattern.pattern
    assert channel_answers._FIGURE_TAG.search("[B1]") and not channel_answers._FIGURE_TAG.search("[S1]")


def test_a_budget_figure_is_numbered_in_an_export_not_left_raw() -> None:
    found = _content()
    assert [(number, figure.label) for number, figure in found.figures] == [(1, "B1")]
    written = " ".join("".join(text for text, _ in block.runs) for block in found.blocks)
    assert "[F1]" in written and not RAW.search(written)


def test_the_csv_holds_every_budget_row_as_a_number() -> None:
    rows = list(csv.DictReader(io.StringIO(csv_bytes(_content()).decode("utf-8-sig"))))
    figures = [row for row in rows if row["section"] == "figure"]
    assert [(row["category"], row["value"]) for row in figures] == [("Total", "48289588"), ("Head Office", "20270110"),
                                                                    ("Public Works", "20232848"), ("Education", "7786630")]
    notes = [row["item"] for row in rows if row["section"] == "note"]
    assert BUDGET_DATA_NOTE in notes and LIVE_DATA_NOTE not in notes  # never "counts of residents' reports"


def test_the_workbook_calls_them_budget_figures_and_charts_them() -> None:
    book = load_workbook(io.BytesIO(xlsx(_content())))
    sheet = book["Figures"]
    cells = [[cell.value for cell in row] for row in sheet.iter_rows()]
    assert sheet["A1"].value == "Budget figures" and any(BUDGET_DATA_NOTE in str(row[0]) for row in cells if row[0])
    assert not any("counts of the reports residents" in str(value) for row in cells for value in row if value)
    assert ["Category", "Amount (GH¢)", "Shown as"] in [row[:3] for row in cells]
    assert ["Head Office", 20270110, "GH¢ 20,270,110"] in [row[:3] for row in cells]
    assert "Chart" in book.sheetnames and len(book["Chart"]._charts) == 1


def test_the_workbook_charts_only_the_bars_the_page_drew() -> None:
    """The page draws the 20 largest and says so; a workbook charting 24 would contradict its own note."""
    drew_two = {**CHART, "categories": CHART["categories"][:2], "series": [{"name": "Figures", "values": CHART["series"][0]["values"][:2]}]}
    book = load_workbook(io.BytesIO(xlsx(_content(drew_two))))
    reference = book["Chart"]._charts[0].series[0].val.numRef.f
    first, last = (int(n) for n in re.findall(r"\$(\d+)", reference))
    assert last - first + 1 == 2


def test_the_pdf_and_word_file_call_them_budget_figures() -> None:
    text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(pdf(_content()))).pages)
    assert "Budget figures" in text and "counts of the reports residents" not in text and not RAW.search(text)
    document = Document(io.BytesIO(docx(_content())))
    body = "\n".join(p.text for p in document.paragraphs)
    assert "Budget figures" in body and BUDGET_DATA_NOTE in body and LIVE_DATA_NOTE not in body and not RAW.search(body)


def test_an_amount_is_a_number_and_a_suppressed_count_never_is() -> None:
    assert figure_value("GH¢ 20,270,110") == 20270110 and figure_value("GHS 800.50") == 800.5 and figure_value("12") == 12
    assert figure_value("none") == 0 and figure_value("fewer than 5") == ""


def _chat_answer() -> dict[str, Any]:
    return {"answer": ANSWER, "status": "answered", "sources": [], "figures": [BUDGET], "chart": None}


def test_a_whatsapp_or_sms_budget_answer_shows_no_raw_citation_and_is_not_called_live_data() -> None:
    chat = channel_answers.for_chat(_chat_answer(), SITE)  # type: ignore[arg-type]
    assert not RAW.search(chat) and channel_answers.BUDGET_DATA_NOTE in chat and channel_answers.LIVE_DATA_NOTE not in chat
    sms = channel_answers.for_sms(_chat_answer(), SITE)  # type: ignore[arg-type]
    assert not RAW.search(" ".join(sms)) and "2026 AMA Budget" in sms[-1]  # the budget it was read from, named


def test_a_spoken_budget_answer_never_reads_a_citation_aloud() -> None:
    spoken = voice_speech.speakable(ANSWER)
    assert "B1" not in spoken and "[" not in spoken


def test_the_translation_check_counts_a_budget_citation_as_a_citation_not_a_figure() -> None:
    numbers, labels = ask_language.figures_and_labels(ANSWER)
    assert labels == {"B1": 1} and "1" not in numbers
    assert not ask_language.survives(ANSWER, ANSWER.replace(" [B1]", ""))  # dropping it is caught as a lost citation
