"""An answer as an Excel workbook: the live figures as cells, a real chart over them, and no sum of a hidden count."""

import io
from typing import Any

import pytest
from openpyxl import load_workbook

from app.services import rag
from app.services.ask_charts import SPREADSHEET_REFUSAL, chart_for
from app.services.ask_export import Answered, content, export_view
from app.services.export_xlsx import xlsx
from tests.test_ask_export import NOW, OPEN, RESOLVED_SMALL, _answer  # the fixtures the other exports are tested with

SITE = "https://nokware.example"


def book(figures: list[dict[str, Any]], chart: dict[str, Any] | None = None,
         question: str = "How many reports by sub-metro?") -> Any:
    # the answer cites both labels: an export numbers only the figures the answer actually cites
    answered = Answered(question, "Okaikoi South has the most [R1][R2].", "answered", [], figures, chart, None)
    return load_workbook(io.BytesIO(xlsx(content(export_view(answered, NOW), SITE))))


def test_a_suppressed_count_is_text_so_a_spreadsheet_cannot_add_it_up() -> None:
    sheet = book([OPEN, RESOLVED_SMALL])["Figures"]
    rows = [(sheet.cell(row=r, column=1).value, sheet.cell(row=r, column=2).value, sheet.cell(row=r, column=3).value)
            for r in range(1, sheet.max_row + 1)]
    assert ("Ashiedu Keteke", 7, "7") in rows  # an exact count is a number in its own column
    suppressed = [row for row in rows if row[2] == "fewer than 5"]
    assert suppressed and all(count is None for _, count, _ in suppressed)  # nothing to total, and nothing to reveal
    assert all(count is None or isinstance(count, int) for _, count, shown in rows if shown and shown != "Shown as")


def test_the_workbook_carries_the_answer_its_sources_and_the_notice() -> None:
    sheet = book([OPEN])["Answer"]
    text = "\n".join(str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value)
    assert "not an official AMA document" in text and "How many reports by sub-metro?" in text
    assert "Okaikoi South has the most" in text


def test_the_chart_is_a_real_excel_chart_over_the_figure_cells() -> None:
    chart, _ = chart_for("reports by sub-metro as a bar chart", [OPEN])
    workbook = book([OPEN], chart)
    drawn = workbook["Chart"]._charts
    assert len(drawn) == 1 and drawn[0].title is not None
    reference = str(drawn[0].series[0].val.numRef.f)
    assert "Figures" in reference and "$B$" in reference  # the cells, not a picture of them


def test_an_answer_without_figures_has_no_chart_sheet() -> None:
    assert "Chart" not in book([]).sheetnames


def test_a_spreadsheet_of_document_figures_is_refused_in_plain_words(monkeypatch: pytest.MonkeyPatch) -> None:
    result = _answer(monkeypatch, "Give me the 2026 budget as a spreadsheet", rag.NO_FIGURES, "Works gets GH¢2m [S1].")
    assert result["answer"].startswith(SPREADSHEET_REFUSAL)
    quiet = _answer(monkeypatch, "What does the 2026 budget give Works?", rag.NO_FIGURES, "Works gets GH¢2m [S1].")
    assert SPREADSHEET_REFUSAL not in quiet["answer"]
